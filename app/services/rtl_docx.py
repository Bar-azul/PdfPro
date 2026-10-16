"""
Hebrew / Arabic support for PDF → Word (pdf2docx).

pdf2docx lays text out left to right. On a right-to-left page that breaks
every line that mixes scripts: MuPDF splits such a line into separate pieces
at each change of direction, pdf2docx writes the pieces in left-to-right page
order, and Word shows "(2026) עם מילה Word" as ") 2026 Word עם מילה (". It
also leaves the paragraphs left to right and gives the runs no complex-script
font or size, so Word falls back to its defaults for the Hebrew letters.

Two passes fix it:

* fix_rtl_layout() runs on pdf2docx's parsed pages, before the .docx is made.
  For each row of a text block that contains RTL letters it takes every
  character with its position, puts them in left-to-right page (visual) order,
  turns that back into reading (logical) order with the Unicode bidi algorithm,
  and rebuilds the row as one line whose spans follow that order. Each
  character keeps the span (font, size, colour, bold) it came from.

* fix_rtl_docx() edits the finished .docx: paragraphs with RTL letters get
  <w:bidi/> (right-to-left paragraph) and runs get <w:rtl/> plus complex-script
  font, size, bold and italic copied from their Latin settings.
"""

import copy
import logging
import re
import zipfile
from pathlib import Path

logger = logging.getLogger(__name__)

_RTL = re.compile(r"[֐-׿؀-ۿݐ-ݿיִ-﷿ﹰ-﻿]")


# ── visual → logical ──────────────────────────────────────────────────────────

def _visual_to_logical_order(text: str) -> list[int]:
    """
    Indexes of `text` (characters in left-to-right page order) in reading order.

    Running the bidi reordering on a visual string of a right-to-left paragraph
    gives back its logical order: right-to-left runs are reversed again, left-to-
    right runs (English, numbers) keep their internal order. Mirroring is not
    applied: PDFs store the logical bracket and draw the mirrored glyph, so the
    character codes are already right for reading order.
    """
    from bidi.algorithm import (get_empty_storage, get_embedding_levels, explicit_embed_and_overrides,
                                resolve_weak_types, resolve_neutral_types, resolve_implicit_levels,
                                reorder_resolved_levels, PARAGRAPH_LEVELS)
    storage = get_empty_storage()
    storage["base_level"] = PARAGRAPH_LEVELS["R"]
    storage["base_dir"] = "R"
    get_embedding_levels(text, storage, False, False)
    for i, ch in enumerate(storage["chars"]):
        ch["_i"] = i
    explicit_embed_and_overrides(storage, False)
    resolve_weak_types(storage, False)
    resolve_neutral_types(storage, False)
    resolve_implicit_levels(storage, False)
    reorder_resolved_levels(storage, False)
    return [ch["_i"] for ch in storage["chars"]]


_NO_SPACE_NEXT_TO = set(".,:;!?")  # glyph gaps next to punctuation are kerning, not spaces
_SWAP_CHARS = "()[]{}"
_SWAP = str.maketrans("()[]{}", ")(][}{")


def _balanced(text: str) -> bool:
    """Brackets open before they close and all close."""
    depth = {"(": 0, "[": 0, "{": 0}
    closes = {")": "(", "]": "[", "}": "{"}
    for c in text:
        if c in depth:
            depth[c] += 1
        elif c in closes:
            depth[closes[c]] -= 1
            if depth[closes[c]] < 0:
                return False
    return all(v == 0 for v in depth.values())


# ── pass 1: parsed layout ─────────────────────────────────────────────────────

def _rows(lines: list) -> list[list]:
    """Group a block's lines into visual rows (lines whose heights overlap)."""
    rows: list[list] = []
    for line in sorted(lines, key=lambda l: (l.bbox[1] + l.bbox[3]) / 2):
        y0, y1 = line.bbox[1], line.bbox[3]
        if rows:
            r0 = min(l.bbox[1] for l in rows[-1])
            r1 = max(l.bbox[3] for l in rows[-1])
            overlap = min(y1, r1) - max(y0, r0)
            if overlap > 0.5 * min(y1 - y0, r1 - r0):
                rows[-1].append(line)
                continue
        rows.append([line])
    return rows


def _fix_row(row: list) -> bool:
    """Rebuild one row in reading order. False when the row is left as it is."""
    from pdf2docx.text.TextSpan import TextSpan

    items = []  # (x0, x1, char, span)
    for line in row:
        if not line.is_horizontal_text:
            return False
        for span in line.spans:
            if not isinstance(span, TextSpan) or not span.chars:
                return False  # images or text without positions: leave the row alone
            for ch in span.chars:
                if len(ch.c) == 1:
                    items.append((ch.bbox[0], ch.bbox[2], ch, span))
                elif ch.c:  # pdf2docx glues a trailing space onto a char ("d "): split it
                    w = (ch.bbox[2] - ch.bbox[0]) / len(ch.c)
                    for k, cp in enumerate(ch.c):
                        part = copy.copy(ch)
                        part.c = cp
                        x = ch.bbox[0] + k * w
                        items.append((x, x + w, part, span))
    text = "".join(c.c for _, _, c, _ in items)
    if not items or not _RTL.search(text):
        return False

    items.sort(key=lambda t: (t[0] + t[1]) / 2)
    # a gap between pieces with no space character in it: add one
    sized = [s.size for _, _, _, s in items if getattr(s, "size", 0)]
    size = sorted(sized)[len(sized) // 2] if sized else 10
    visual = []
    prev_x1 = None
    for x0, x1, ch, span in items:
        if (prev_x1 is not None and x0 - prev_x1 > 0.2 * size
                and ch.c != " " and visual and visual[-1][0].c != " "
                and ch.c not in _NO_SPACE_NEXT_TO and visual[-1][0].c not in _NO_SPACE_NEXT_TO):
            space = copy.copy(ch)
            space.c = " "
            visual.append((space, visual[-1][1]))
        visual.append((ch, span))
        prev_x1 = x1

    # a "character" can hold more than one code point (ligatures, marks): order by
    # code point, then map each back to its item, keeping each item once
    owner = [k for k, (c, _) in enumerate(visual) for _ in c.c]
    seen, logical = set(), []
    for i in _visual_to_logical_order("".join(c.c for c, _ in visual)):
        k = owner[i]
        if k not in seen:
            seen.add(k)
            logical.append(visual[k])

    # consecutive characters from the same span become one span, in reading order
    # Some PDF producers store the bracket character as drawn (mirrored) rather
    # than as read, which leaves ")2026(" in reading order. If the brackets only
    # balance when swapped, swap them.
    text = "".join(c.c for c, _ in logical)
    if not _balanced(text) and _balanced(text.translate(_SWAP)):
        swapped = []
        for ch, span in logical:
            if ch.c in _SWAP_CHARS:
                ch = copy.copy(ch)
                ch.c = ch.c.translate(_SWAP)
            swapped.append((ch, span))
        logical = swapped

    # pieces split at a change of direction often both carry the space between them
    collapsed = []
    for ch, span in logical:
        if ch.c == " " and (not collapsed or collapsed[-1][0].c == " "):
            continue
        collapsed.append((ch, span))
    new_spans = []
    for ch, span in collapsed:
        if new_spans and new_spans[-1][0] is span:
            new_spans[-1][1].append(ch)
        else:
            new_spans.append((span, [ch]))
    rebuilt = []
    for span, chars in new_spans:
        s = copy.copy(span)
        s.chars = chars
        rebuilt.append(s)

    first = row[0]
    x0 = min(l.bbox[0] for l in row); y0 = min(l.bbox[1] for l in row)
    x1 = max(l.bbox[2] for l in row); y1 = max(l.bbox[3] for l in row)
    first.spans.reset(rebuilt)
    first.update_bbox((x0, y0, x1, y1))
    first.line_break = any(getattr(l, "line_break", 0) for l in row)
    first.tab_stop = 0
    for line in row[1:]:
        line.spans.reset([])
    return True


def fix_rtl_layout(converter) -> int:
    """Put RTL rows of every parsed page in reading order. Returns rows changed."""
    from pdf2docx.text.TextBlock import TextBlock
    changed = 0

    def walk_blocks(blocks):
        nonlocal changed
        for blk in blocks:
            if isinstance(blk, TextBlock):
                lines = list(blk.lines)
                if not _RTL.search("".join(getattr(s, "text", "") or "" for l in lines for s in l.spans)):
                    continue
                for row in _rows(lines):
                    try:
                        changed += _fix_row(row)
                    except Exception as e:  # never fail the conversion over one row
                        logger.debug(f"RTL row left as is: {e}")
                blk.lines.reset([l for l in lines if len(l.spans)])
            elif getattr(blk, "is_table_block", False):  # tables: cells hold their own blocks
                for row in blk:
                    for cell in row:
                        inner = getattr(cell, "blocks", None)
                        if inner is not None:
                            walk_blocks(inner)

    for page in converter.pages:
        if not getattr(page, "finalized", True):
            continue
        for section in page.sections:
            for column in section:
                walk_blocks(column.blocks)
    return changed


# ── pass 2: the .docx ─────────────────────────────────────────────────────────

_W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"


def fix_rtl_docx(path: Path) -> int:
    """Mark RTL paragraphs and runs in a .docx written by pdf2docx. Returns paragraphs changed."""
    from docx import Document
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    doc = Document(str(path))
    changed = 0

    def paragraphs(container):
        for p in container.paragraphs:
            yield p
        for table in container.tables:
            for row in table.rows:
                for cell in row.cells:
                    yield from paragraphs(cell)

    def first_child(parent, tag):
        return parent.find(qn(tag))

    for p in paragraphs(doc):
        if not _RTL.search(p.text or ""):
            continue
        pPr = p._p.get_or_add_pPr()
        if first_child(pPr, "w:bidi") is None:
            bidi = OxmlElement("w:bidi")
            # w:bidi must come before w:spacing / w:ind / w:jc in pPr
            anchor = None
            for tag in ("w:adjustRightInd", "w:snapToGrid", "w:spacing", "w:ind",
                        "w:contextualSpacing", "w:mirrorIndents", "w:suppressOverlap", "w:jc",
                        "w:textDirection", "w:textAlignment", "w:outlineLvl", "w:rPr", "w:sectPr"):
                anchor = first_child(pPr, tag)
                if anchor is not None:
                    break
            if anchor is not None:
                anchor.addprevious(bidi)
            else:
                pPr.append(bidi)
        # In a right-to-left paragraph Word reads jc left/right as start/end, so
        # pdf2docx's "right" (the page's right edge) is "left" here, and vice versa.
        jc = first_child(pPr, "w:jc")
        if jc is not None and jc.get(qn("w:val")) in ("left", "right"):
            jc.set(qn("w:val"), "left" if jc.get(qn("w:val")) == "right" else "right")
        ind = first_child(pPr, "w:ind")
        if ind is not None:  # indents are start/end too
            left, right = ind.get(qn("w:left")), ind.get(qn("w:right"))
            if left is not None or right is not None:
                ind.set(qn("w:left"), right or "0")
                ind.set(qn("w:right"), left or "0")

        for r in p.runs:
            rPr = r._r.get_or_add_rPr()
            if _RTL.search(r.text or "") and first_child(rPr, "w:rtl") is None:
                rtl = OxmlElement("w:rtl")  # schema order: before cs/em/lang/...
                after = next((first_child(rPr, t) for t in ("w:cs", "w:em", "w:lang", "w:eastAsianLayout",
                                                           "w:specVanish", "w:oMath")
                              if first_child(rPr, t) is not None), None)
                if after is not None:
                    after.addprevious(rtl)
                else:
                    rPr.append(rtl)
            fonts = first_child(rPr, "w:rFonts")
            if fonts is not None and fonts.get(qn("w:cs")) is None and fonts.get(qn("w:ascii")):
                fonts.set(qn("w:cs"), fonts.get(qn("w:ascii")))
            for latin, cs in (("w:sz", "w:szCs"), ("w:b", "w:bCs"), ("w:i", "w:iCs")):
                src = first_child(rPr, latin)
                if src is not None and first_child(rPr, cs) is None:
                    el = OxmlElement(cs)
                    if src.get(qn("w:val")) is not None:
                        el.set(qn("w:val"), src.get(qn("w:val")))
                    src.addnext(el)
        changed += 1

    if changed:
        doc.save(str(path))
    return changed
